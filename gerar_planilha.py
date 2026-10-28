"""
Gera Controle_Alugueis_Demo.xlsx: controle de aluguéis para administração de imóveis
no Brasil (R$, multa e juros de mora, taxa de administração, repasse aos proprietários,
reajuste anual). Todos os nomes, endereços e valores são fictícios.
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Controle_Alugueis_Demo.xlsx"

REF = dt.date(2026, 10, 28)       # data de referência
MULTA = 0.10                      # multa sobre o aluguel em atraso
JUROS = 0.01                      # juros de mora ao mês, pro rata die
TOL = 0                           # dias de tolerância
TAXA_ADM = 0.10                   # taxa de administração sobre o valor recebido

NAVY = "1F3864"
BRL = '[$-416]"R$" #,##0.00'
DATA = "[$-416]dd/mm/yyyy"
MES = "[$-416]mmm/yy"
PCT = "[$-416]0.0%"

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FONT = Font(color="0000FF")
LINK_FONT = Font(color="008000")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ---------------------------------------------------------------- dados fictícios
IMOVEIS = [
    ("I01", "Rua Exemplo A, 100, ap. 32", "Batel", "Apartamento"),
    ("I02", "Rua Exemplo B, 250", "Água Verde", "Casa"),
    ("I03", "Rua Exemplo C, 45, sala 4", "Centro", "Sala comercial"),
    ("I04", "Rua Exemplo D, 780, ap. 101", "Bigorrilho", "Apartamento"),
    ("I05", "Rua Exemplo E, 15", "Portão", "Casa"),
]
CONTRATOS = [  # contrato, imóvel, inquilino, início, aluguel, dia venc., índice, garantia
    ("CT-01", "I01", "Marina L.", dt.date(2025, 11, 15), 2800.00, 5, "IGP-M", "Seguro fiança"),
    ("CT-02", "I02", "Pedro R.", dt.date(2024, 3, 1), 3500.00, 5, "IPCA", "Fiador"),
    ("CT-03", "I03", "Lucas F. ME", dt.date(2026, 1, 10), 1900.00, 10, "IGP-M", "Caução (3 aluguéis)"),
    ("CT-04", "I04", "Juliana T.", dt.date(2025, 12, 5), 4200.00, 5, "IPCA", "Seguro fiança"),
    ("CT-05", "I05", "Rafael G.", dt.date(2025, 6, 1), 2300.00, 5, "IGP-M", "Fiador"),
]
PROPRIETARIOS = [  # nome, % por imóvel I01..I05
    ("Ana P.", [1.0, 0.5, 0.0, 0.0, 1.0]),
    ("Bruno S.", [0.0, 0.5, 0.0, 0.6, 0.0]),
    ("Carlos M.", [0.0, 0.0, 1.0, 0.4, 0.0]),
]
# custos da cobrança judicial: data, imóvel, descrição, valor, quem pagou, reembolsável pelo inquilino
CUSTOS_JUD = [
    (dt.date(2026, 10, 15), "I05", "Custas iniciais da ação de despejo por falta de pagamento", 480.00, "Proprietário", "Sim"),
    (dt.date(2026, 10, 15), "I05", "Honorários advocatícios contratuais (entrada)", 1500.00, "Proprietário", "Não"),
    (dt.date(2026, 10, 22), "I05", "Diligência do oficial de justiça (citação)", 95.00, "Proprietário", "Sim"),
]
# tratativa manual por cobrança: (imóvel, mês) -> (tratativa, observação)
TRATATIVAS = {
    ("I03", 10): ("Em acordo", "Acordo em 15/10: restante com multa e juros até 10/11/2026"),
    ("I05", 8): ("Cobrança judicial", "Ação de despejo por falta de pagamento ajuizada em 15/10/2026"),
    ("I05", 9): ("Cobrança judicial", "Incluída na ação de despejo"),
    ("I05", 10): ("Cobrança judicial", "Incluída na ação de despejo"),
}
DESPESAS = [  # data, imóvel, descrição, valor (pagas pela administradora e descontadas do proprietário)
    (dt.date(2026, 7, 20), "I04", "Seguro incêndio anual", 380.00),
    (dt.date(2026, 8, 12), "I02", "Reparo hidráulico", 450.00),
    (dt.date(2026, 9, 18), "I03", "Pintura da fachada (parte do proprietário)", 1200.00),
]
MESES = [dt.date(2026, m, 1) for m in range(7, 13)]


def venc(comp, dia):
    return dt.date(comp.year, comp.month, dia)


def total_quitacao(aluguel, dias):
    multa = round(aluguel * MULTA, 2) if dias > TOL else 0
    juros = round(aluguel * JUROS / 30 * dias, 2) if dias > TOL else 0
    return round(aluguel + multa + juros, 2)


# recebimentos: (data, cobrança, valor, forma, observação)
cobrancas = []  # (id, contrato, imóvel, competência)
n = 1
for ct in CONTRATOS:
    for comp in MESES:
        cobrancas.append((f"C-{n:02d}", ct[0], ct[1], comp))
        n += 1

recebimentos = []
for cid, ctid, im, comp in cobrancas:
    ct = next(c for c in CONTRATOS if c[0] == ctid)
    v = venc(comp, ct[5])
    if v > REF:
        continue
    aluguel = ct[4]
    forma = "Pix" if im in ("I01", "I03") else "Boleto"
    if im == "I02" and comp.month == 10:          # pago 8 dias atrasado, com multa e juros
        d = v + dt.timedelta(days=8)
        recebimentos.append((d, cid, total_quitacao(aluguel, 8), forma, "Pago com atraso, multa e juros incluídos"))
    elif im == "I04" and comp.month == 9:         # pago 2 dias atrasado
        d = v + dt.timedelta(days=2)
        recebimentos.append((d, cid, total_quitacao(aluguel, 2), forma, "Pago com atraso, multa e juros incluídos"))
    elif im == "I03" and comp.month == 10:        # pagamento parcial
        recebimentos.append((v, cid, 1000.00, forma, "Pagamento parcial; inquilino prometeu o restante"))
    elif im == "I04" and comp.month == 10:        # em aberto
        continue
    elif im == "I05" and comp.month >= 8:         # sem pagamento desde agosto: cobrança judicial
        continue
    else:
        recebimentos.append((v, cid, aluguel, forma, ""))
recebimentos.sort(key=lambda r: r[0])

# repasses de julho a setembro, calculados com as mesmas regras da planilha
idx = {im[0]: i for i, im in enumerate(IMOVEIS)}
repasses = []
for comp in MESES[:3]:
    liquido_im = [0.0] * len(IMOVEIS)
    for d, cid, valor, _, _ in recebimentos:
        if d.year == comp.year and d.month == comp.month:
            im = next(c[2] for c in cobrancas if c[0] == cid)
            liquido_im[idx[im]] += valor * (1 - TAXA_ADM)
    for d, im, _, valor in DESPESAS:
        if d.month == comp.month:
            liquido_im[idx[im]] -= valor
    data_rep = dt.date(comp.year, comp.month + 1, 15)
    for nome, pcts in PROPRIETARIOS:
        v = round(sum(p * l for p, l in zip(pcts, liquido_im)), 2)
        if v:
            repasses.append((data_rep, nome, v, f"Repasse de {comp.strftime('%m/%Y')}"))


# ---------------------------------------------------------------- helpers
def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row, i, lab)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, color="595959")


def box(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.border = BOX


wb = Workbook()
ws_guia = wb.active
ws_guia.title = "Guia"
ws_p = wb.create_sheet("Painel")
ws_e = wb.create_sheet("Extrato")
ws_r = wb.create_sheet("Recebimentos")
ws_c = wb.create_sheet("Cobranças")
ws_ct = wb.create_sheet("Contratos")
ws_i = wb.create_sheet("Imóveis")
ws_d = wb.create_sheet("Despesas")
ws_pr = wb.create_sheet("Proprietários")
ws_rp = wb.create_sheet("Repasses")
ws_par = wb.create_sheet("Parâmetros")

P = "'Parâmetros'!"
REFC, MULC, JURC, TOLC, ADMC = (P + "$B$4", P + "$B$5", P + "$B$6", P + "$B$7", P + "$B$8")

# ---------------------------------------------------------------- Parâmetros
title(ws_par, "Parâmetros", "Células amarelas podem ser alteradas; todas as abas recalculam.")
rows = [
    ("Data de referência", REF, DATA, "Data usada para calcular atrasos e o que já venceu"),
    ("Multa por atraso", MULTA, "0%", "Sobre o valor do aluguel, conforme contrato"),
    ("Juros de mora ao mês", JUROS, "0.0%", "Pro rata die: juros/30 por dia de atraso"),
    ("Dias de tolerância", TOL, "0", "Atraso até este número de dias não gera multa nem juros"),
    ("Taxa de administração", TAXA_ADM, "0%", "Sobre o valor recebido; descontada no repasse"),
]
header(ws_par, 3, ["Parâmetro", "Valor", "Observação"], [28, 16, 60])
for i, (lab, val, fmt, obs) in enumerate(rows, 4):
    ws_par.cell(i, 1, lab)
    c = ws_par.cell(i, 2, val)
    c.number_format, c.fill, c.font = fmt, YELLOW, INPUT_FONT
    ws_par.cell(i, 3, obs)
box(ws_par, "A4:C8")

ws_j = wb.create_sheet("Cobrança judicial", index=wb.sheetnames.index("Despesas") + 1)
NI = len(IMOVEIS)
LI = 4 + NI                       # última linha de Imóveis / Contratos
NP = len(PROPRIETARIOS)
PCOLS = [chr(ord("H") + j) for j in range(NI)]   # colunas de participação em Proprietários
PC1, PCN = PCOLS[0], PCOLS[-1]
SIT_CORES = [("Em atraso", "F8CBAD"), ("Parcial", "FCE4D6"), ("Em acordo", "DDEBF7"),
             ("Cobrança judicial", "D9D2E9"), ("Pago com atraso", "FFF2CC"), ("Pago", "E2EFDA")]


def cores_situacao(ws, rng, first_cell):
    for txt, color in SIT_CORES:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first_cell}="{txt}"'],
                                                       fill=PatternFill("solid", fgColor=color)))


# ---------------------------------------------------------------- Imóveis
title(ws_i, "Imóveis", "Curitiba - PR. Endereços fictícios.")
header(ws_i, 4, ["Imóvel", "Endereço", "Bairro", "Tipo", "Contrato"], [9, 34, 14, 16, 10])
for r, (iid, end, bairro, tipo) in enumerate(IMOVEIS, 5):
    for col, v in enumerate([iid, end, bairro, tipo], 1):
        ws_i.cell(r, col, v).font = INPUT_FONT
    ws_i.cell(r, 5, f'=IFERROR(INDEX(Contratos!$A$5:$A${LI},MATCH(A{r},Contratos!$B$5:$B${LI},0)),"Vago")').font = LINK_FONT
box(ws_i, f"A5:E{LI}")

# ---------------------------------------------------------------- Contratos
title(ws_ct, "Contratos de locação", "Valor do aluguel, vencimento, índice de reajuste e garantia.")
header(ws_ct, 4, ["Contrato", "Imóvel", "Inquilino", "Início", "Aluguel (R$)", "Dia do vencimento",
                  "Índice de reajuste", "Garantia", "Próximo reajuste", "Dias até o reajuste"],
       [10, 8, 16, 12, 14, 11, 11, 20, 14, 11])
for r, ct in enumerate(CONTRATOS, 5):
    for col, v in enumerate(ct, 1):
        ws_ct.cell(r, col, v).font = INPUT_FONT
    ws_ct.cell(r, 4).number_format = DATA
    ws_ct.cell(r, 5).number_format = BRL
    ws_ct.cell(r, 9, f"=DATE(YEAR({REFC})+IF(DATE(YEAR({REFC}),MONTH(D{r}),DAY(D{r}))<={REFC},1,0),MONTH(D{r}),DAY(D{r}))").number_format = DATA
    ws_ct.cell(r, 10, f"=I{r}-{REFC}")
box(ws_ct, f"A5:J{LI}")
ws_ct.conditional_formatting.add(f"I5:J{LI}", FormulaRule(formula=["$J5<=60"], fill=PatternFill("solid", fgColor="FCE4D6")))

# ---------------------------------------------------------------- Recebimentos
title(ws_r, "Recebimentos", "Lance aqui cada pagamento recebido. Use o código da cobrança (aba Cobranças).")
header(ws_r, 4, ["Data", "Cobrança", "Valor (R$)", "Forma", "Observação"], [12, 10, 14, 10, 48])
for r, (d, cid, valor, forma, obs) in enumerate(recebimentos, 5):
    for col, v in enumerate([d, cid, valor, forma, obs], 1):
        ws_r.cell(r, col, v).font = INPUT_FONT
    ws_r.cell(r, 1).number_format = DATA
    ws_r.cell(r, 3).number_format = BRL
box(ws_r, f"A5:E{4 + len(recebimentos)}")
RR = "Recebimentos!$B$5:$B$200"

# ---------------------------------------------------------------- Cobranças
title(ws_c, "Cobranças mensais",
      "Uma linha por contrato e mês. Multa, juros, saldo, faixa de atraso e situação são calculados. "
      "A coluna Tratativa (lista) marca acordos e cobranças judiciais.")
header(ws_c, 4, ["Cobrança", "Contrato", "Imóvel", "Inquilino", "Competência", "Vencimento", "Aluguel (R$)",
                 "Recebido (R$)", "Último pagamento", "Dias de atraso", "Multa (R$)", "Juros (R$)",
                 "Total devido (R$)", "Saldo em aberto (R$)", "Situação", "Aluguel recebido (R$)",
                 "Tratativa", "Faixa de atraso", "Observação"],
       [10, 9, 8, 14, 13, 12, 13, 13, 12, 9, 11, 10, 13, 13, 17, 13, 17, 14, 52])
CT = "Contratos!"
for r, (cid, ctid, im, comp) in enumerate(cobrancas, 5):
    look = lambda col: f"INDEX({CT}${col}$5:${col}${LI},MATCH(B{r},{CT}$A$5:$A${LI},0))"
    ws_c.cell(r, 1, cid)
    ws_c.cell(r, 2, ctid).font = INPUT_FONT
    ws_c.cell(r, 3, f"={look('B')}").font = LINK_FONT
    ws_c.cell(r, 4, f"={look('C')}").font = LINK_FONT
    c = ws_c.cell(r, 5, comp); c.number_format = MES; c.font = INPUT_FONT
    ws_c.cell(r, 6, f"=DATE(YEAR(E{r}),MONTH(E{r}),{look('F')})").number_format = DATA
    ws_c.cell(r, 7, f"={look('E')}").number_format = BRL
    ws_c.cell(r, 8, f"=SUMIFS(Recebimentos!$C$5:$C$200,{RR},A{r})").number_format = BRL
    ws_c.cell(r, 9, f'=IF(H{r}=0,"",_xlfn.MAXIFS(Recebimentos!$A$5:$A$200,{RR},A{r}))').number_format = DATA
    quitado = (f"IF(I{r}=\"\",FALSE,H{r}>=G{r}*(1+IF(MAX(0,I{r}-F{r})>{TOLC},{MULC},0))"
               f"+IF(MAX(0,I{r}-F{r})>{TOLC},G{r}*{JURC}/30*MAX(0,I{r}-F{r}),0)-0.01)")
    ws_c.cell(r, 10, f"=IF(F{r}>{REFC},0,IF({quitado},MAX(0,I{r}-F{r}),MAX(0,{REFC}-F{r})))")
    ws_c.cell(r, 11, f"=IF(J{r}>{TOLC},ROUND(G{r}*{MULC},2),0)").number_format = BRL
    ws_c.cell(r, 12, f"=IF(J{r}>{TOLC},ROUND(G{r}*{JURC}/30*J{r},2),0)").number_format = BRL
    ws_c.cell(r, 13, f"=IF(F{r}>{REFC},0,G{r}+K{r}+L{r})").number_format = BRL
    ws_c.cell(r, 14, f"=MAX(0,ROUND(M{r}-H{r},2))").number_format = BRL
    ws_c.cell(r, 15, f'=IF(F{r}>{REFC},"A vencer",IF(N{r}=0,IF(J{r}>{TOLC},"Pago com atraso","Pago"),'
                     f'IF(Q{r}<>"",Q{r},IF(H{r}>0,"Parcial","Em atraso"))))')
    ws_c.cell(r, 16, f"=IF(F{r}>{REFC},0,MIN(G{r},H{r}))").number_format = BRL
    trat, obs = TRATATIVAS.get((im, comp.month), ("", ""))
    ws_c.cell(r, 17, trat or None).font = INPUT_FONT
    ws_c.cell(r, 17).fill = YELLOW
    ws_c.cell(r, 18, f'=IF(N{r}=0,"",IF(J{r}<=30,"1 a 30 dias",IF(J{r}<=60,"31 a 60 dias",'
                     f'IF(J{r}<=90,"61 a 90 dias","Mais de 90 dias"))))')
    ws_c.cell(r, 19, obs or None).font = INPUT_FONT
LASTC = 4 + len(cobrancas)
box(ws_c, f"A5:S{LASTC}")
cores_situacao(ws_c, f"O5:O{LASTC}", "$O5")
dv_t = DataValidation(type="list", formula1='"Em acordo,Cobrança judicial"', allow_blank=True)
ws_c.add_data_validation(dv_t)
dv_t.add(f"Q5:Q{LASTC}")
ws_c.freeze_panes = "B5"

C = "'Cobranças'!"
CR = lambda col: f"{C}${col}$5:${col}${LASTC}"

# ---------------------------------------------------------------- Despesas
title(ws_d, "Despesas do proprietário", "Pagas pela administradora e descontadas no repasse.")
header(ws_d, 4, ["Data", "Imóvel", "Descrição", "Valor (R$)"], [12, 8, 44, 13])
for r, (d, im, desc, v) in enumerate(DESPESAS, 5):
    for col, val in enumerate([d, im, desc, v], 1):
        ws_d.cell(r, col, val).font = INPUT_FONT
    ws_d.cell(r, 1).number_format = DATA
    ws_d.cell(r, 4).number_format = BRL
box(ws_d, f"A5:D{4 + len(DESPESAS)}")

# ---------------------------------------------------------------- Cobrança judicial
title(ws_j, "Custos da cobrança judicial",
      "Custas, honorários e diligências de cada processo. Os pagos pelo proprietário são descontados no repasse.")
header(ws_j, 4, ["Data", "Imóvel", "Contrato", "Descrição", "Valor (R$)", "Pago por", "Reembolsável pelo inquilino?"],
       [12, 8, 10, 52, 13, 14, 16])
for r, (d, im, desc, v, quem, reemb) in enumerate(CUSTOS_JUD, 5):
    for col, val in [(1, d), (2, im), (4, desc), (5, v), (6, quem), (7, reemb)]:
        ws_j.cell(r, col, val).font = INPUT_FONT
    ws_j.cell(r, 3, f"=IFERROR(INDEX({CT}$A$5:$A${LI},MATCH(B{r},{CT}$B$5:$B${LI},0)),\"\")").font = LINK_FONT
    ws_j.cell(r, 1).number_format = DATA
    ws_j.cell(r, 5).number_format = BRL
LJ = 4 + len(CUSTOS_JUD)
box(ws_j, f"A5:G{LJ}")
ws_j.cell(LJ + 2, 4, "Total").font = Font(bold=True)
c = ws_j.cell(LJ + 2, 5, "=SUM(E5:E200)" if False else f"=SUM(E5:E{LJ + 1})")
c.number_format, c.font = BRL, Font(bold=True)
ws_j.cell(LJ + 3, 4, "Reembolsável pelo inquilino (pedir na ação)")
c = ws_j.cell(LJ + 3, 5, f'=SUMIFS(E5:E{LJ + 1},G5:G{LJ + 1},"Sim")')
c.number_format = BRL
ws_j.cell(LJ + 5, 1, "Os valores reembolsáveis dependem da decisão judicial e do contrato; confirme com o advogado responsável.").font = Font(italic=True, color="595959")
dv_q = DataValidation(type="list", formula1='"Proprietário,Administradora"', allow_blank=True)
dv_s = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
ws_j.add_data_validation(dv_q); ws_j.add_data_validation(dv_s)
dv_q.add("F5:F200"); dv_s.add("G5:G200")
JUD = lambda col: f"'Cobrança judicial'!${col}$5:${col}$200"

# ---------------------------------------------------------------- Painel
title(ws_p, "Painel da carteira de aluguéis", "Exemplo com dados fictícios. Fórmulas apenas, sem macros. Abre no Excel e no Google Sheets.")
ws_p["A3"] = "Data de referência:"
ws_p["B3"] = f"={REFC}"
ws_p["B3"].number_format = DATA
ws_p["B3"].font = LINK_FONT
ws_p["A5"] = "Indicadores"
ws_p["A5"].font = Font(bold=True)
kpis = [
    ("Aluguel vencido até a data", f"=SUMIFS({CR('G')},{CR('F')},\"<=\"&{REFC})", BRL, NAVY),
    ("Aluguel recebido (sem multa e juros)", f"=SUM({CR('P')})", BRL, NAVY),
    ("Taxa de recebimento", "=IF(B6=0,0,B7/B6)", PCT, NAVY),
    ("Em aberto (com multa e juros)", f"=SUM({CR('N')})", BRL, "C00000"),
    ("Cobranças em atraso ou parciais", f"=COUNTIF({CR('O')},\"Em atraso\")+COUNTIF({CR('O')},\"Parcial\")", "0", "C00000"),
    ("Em acordo (saldo)", f"=SUMIFS({CR('N')},{CR('O')},\"Em acordo\")", BRL, NAVY),
    ("Em cobrança judicial (saldo)", f"=SUMIFS({CR('N')},{CR('O')},\"Cobrança judicial\")", BRL, "C00000"),
    ("Custos da cobrança judicial", f"=SUMIFS({JUD('E')},{JUD('A')},\"<=\"&{REFC})", BRL, "C00000"),
    ("Reajustes nos próximos 60 dias", f"=COUNTIFS(Contratos!$J$5:$J${LI},\"<=60\")", "0", NAVY),
    ("A repassar aos proprietários", f"='Proprietários'!F{5 + NP + 1}", BRL, NAVY),
]
for i, (lab, f, fmt, color) in enumerate(kpis, 6):
    ws_p.cell(i, 1, lab).font = Font(bold=True)
    c = ws_p.cell(i, 2, f)
    c.number_format, c.font = fmt, Font(bold=True, size=12, color=color)
    for col in (1, 2):
        ws_p.cell(i, col).fill = PatternFill("solid", fgColor="EEF3F8")
        ws_p.cell(i, col).border = BOX

# atraso por faixa (aging), ao lado dos indicadores
ws_p["D5"] = "Atraso por faixa (cobranças com saldo)"
ws_p["D5"].font = Font(bold=True)
for j, lab in enumerate(["Faixa", "Cobranças", "Saldo (R$)"]):
    c = ws_p.cell(6, 4 + j, lab)
    c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
    c.alignment = Alignment(horizontal="center")
for k, faixa in enumerate(["1 a 30 dias", "31 a 60 dias", "61 a 90 dias", "Mais de 90 dias"]):
    r = 7 + k
    ws_p.cell(r, 4, faixa)
    ws_p.cell(r, 5, f"=COUNTIF({CR('R')},D{r})")
    ws_p.cell(r, 6, f"=SUMIFS({CR('N')},{CR('R')},D{r})").number_format = BRL
ws_p.cell(11, 4, "Total").font = Font(bold=True)
ws_p.cell(11, 5, "=SUM(E7:E10)").font = Font(bold=True)
c = ws_p.cell(11, 6, "=SUM(F7:F10)"); c.number_format, c.font = BRL, Font(bold=True)
box(ws_p, "D7:F11")

TOP = 18                           # cabeçalho da tabela por imóvel
R1, RN = TOP + 1, TOP + NI          # linhas dos imóveis
RT = RN + 1                         # total
ws_p.cell(TOP - 1, 1, "Por imóvel").font = Font(bold=True)
header(ws_p, TOP, ["Imóvel", "Endereço", "Aluguel vencido (R$)", "Recebido com encargos (R$)", "Em aberto (R$)",
                   "Taxa de adm. (R$)", "Despesas (R$)", "Custos judiciais (R$)", "Líquido ao proprietário (R$)",
                   "Taxa de recebimento", "Situação"],
       [36, 34, 15, 15, 14, 13, 13, 14, 17, 13, 17])
ws_p.row_dimensions[TOP].height = 46
for k in range(NI):
    r, ir = R1 + k, 5 + k
    ws_p.cell(r, 1, f"='Imóveis'!A{ir}").font = LINK_FONT
    ws_p.cell(r, 2, f"='Imóveis'!B{ir}&\" - \"&'Imóveis'!C{ir}").font = LINK_FONT
    ws_p.cell(r, 3, f"=SUMIFS({CR('G')},{CR('C')},A{r},{CR('F')},\"<=\"&{REFC})").number_format = BRL
    ws_p.cell(r, 4, f"=SUMIFS({CR('H')},{CR('C')},A{r})").number_format = BRL
    ws_p.cell(r, 5, f"=SUMIFS({CR('N')},{CR('C')},A{r})").number_format = BRL
    ws_p.cell(r, 6, f"=ROUND(D{r}*{ADMC},2)").number_format = BRL
    ws_p.cell(r, 7, f"=SUMIFS(Despesas!$D$5:$D$200,Despesas!$B$5:$B$200,A{r},Despesas!$A$5:$A$200,\"<=\"&{REFC})").number_format = BRL
    ws_p.cell(r, 8, f"=SUMIFS({JUD('E')},{JUD('B')},A{r},{JUD('F')},\"Proprietário\",{JUD('A')},\"<=\"&{REFC})").number_format = BRL
    ws_p.cell(r, 9, f"=D{r}-F{r}-G{r}-H{r}").number_format = BRL
    ws_p.cell(r, 10, f"=IF(C{r}=0,0,SUMIFS({CR('P')},{CR('C')},A{r})/C{r})").number_format = PCT
    ws_p.cell(r, 11, f'=IF(COUNTIFS({CR("C")},A{r},{CR("O")},"Cobrança judicial")>0,"Cobrança judicial",'
                     f'IF(COUNTIFS({CR("C")},A{r},{CR("O")},"Em acordo")>0,"Em acordo",IF(E{r}>0,"Em atraso","Em dia")))')
ws_p.cell(RT, 1, "Total").font = Font(bold=True)
for col in "CDEFGHI":
    c = ws_p[f"{col}{RT}"]
    c.value = f"=SUM({col}{R1}:{col}{RN})"
    c.number_format, c.font = BRL, Font(bold=True)
ws_p[f"J{RT}"] = f"=IF(C{RT}=0,0,B7/C{RT})"
ws_p[f"J{RT}"].number_format, ws_p[f"J{RT}"].font = PCT, Font(bold=True)
box(ws_p, f"A{R1}:K{RT}")
cores_situacao(ws_p, f"K{R1}:K{RN}", f"$K{R1}")
ws_p.conditional_formatting.add(f"K{R1}:K{RN}", FormulaRule(formula=[f'$K{R1}="Em dia"'], fill=PatternFill("solid", fgColor="E2EFDA")))

PO = RT + 2                         # seção de proprietários
ws_p.cell(PO, 1, "Proprietários").font = Font(bold=True)
for i, lab in enumerate(["Proprietário", "Líquido acumulado (R$)", "Já repassado (R$)", "A repassar (R$)"], 1):
    c = ws_p.cell(PO + 1, i, lab)
    c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
    c.alignment = Alignment(horizontal="center", wrap_text=True)
for k in range(NP):
    r, pr = PO + 2 + k, 5 + k
    ws_p.cell(r, 1, f"='Proprietários'!A{pr}").font = LINK_FONT
    for col, src in zip((2, 3, 4), ("D", "E", "F")):
        c = ws_p.cell(r, col, f"='Proprietários'!{src}{pr}")
        c.number_format, c.font = BRL, LINK_FONT
box(ws_p, f"A{PO + 2}:D{PO + 1 + NP}")

ch = BarChart()
ch.type = "col"
ch.title = "Recebido x em aberto por imóvel"
ch.add_data(Reference(ws_p, min_col=4, max_col=5, min_row=TOP, max_row=RN), titles_from_data=True)
ch.set_categories(Reference(ws_p, min_col=1, min_row=R1, max_row=RN))
ch.y_axis.number_format = '[$-416]"R$" #,##0'
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.title.overlay = False
ch.legend.position = "b"
ch.legend.overlay = False
ch.height, ch.width = 8, 17
ws_p.add_chart(ch, f"F{PO}")

# ---------------------------------------------------------------- Proprietários
title(ws_pr, "Proprietários e participação", "Percentual de cada proprietário em cada imóvel. O líquido segue a participação.")
header(ws_pr, 4, ["Proprietário", "Participação por imóvel", "", "Líquido acumulado (R$)",
                  "Já repassado (R$)", "A repassar (R$)"], [14, 14, 14, 15, 14, 14])
ws_pr.merge_cells("B4:C4")
ws_pr["H3"] = "Participação (%)"
ws_pr["H3"].font = Font(bold=True)
for j, (iid, *_rest) in enumerate(IMOVEIS):
    c = ws_pr.cell(4, 8 + j, iid)
    c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BOX
LIQ = 5 + NP + 3                    # linha do líquido por imóvel
for k, (nome, pcts) in enumerate(PROPRIETARIOS):
    r = 5 + k
    ws_pr.cell(r, 1, nome).font = INPUT_FONT
    for j, p in enumerate(pcts):
        c = ws_pr.cell(r, 8 + j, p)
        c.number_format, c.font, c.border = "0%", INPUT_FONT, BOX
    ws_pr.cell(r, 2, "=" + "&\" / \"&".join(f'TEXT({col}{r},"0%")' for col in PCOLS))
    ws_pr.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws_pr.cell(r, 4, f"=SUMPRODUCT({PC1}{r}:{PCN}{r},${PC1}${LIQ}:${PCN}${LIQ})").number_format = BRL
    ws_pr.cell(r, 5, f"=SUMIFS(Repasses!$C$5:$C$200,Repasses!$B$5:$B$200,A{r})").number_format = BRL
    ws_pr.cell(r, 6, f"=ROUND(D{r}-E{r},2)").number_format = BRL
ws_pr.cell(5 + NP, 1, "Soma da participação (deve ser 100%)")
for col in PCOLS:
    c = ws_pr[f"{col}{5 + NP}"]
    c.value = f"=SUM({col}5:{col}{4 + NP})"
    c.number_format = "0%"
ws_pr.cell(5 + NP + 1, 1, "Total").font = Font(bold=True)
for col in "DEF":
    c = ws_pr[f"{col}{5 + NP + 1}"]
    c.value = f"=SUM({col}5:{col}{4 + NP})"
    c.number_format, c.font = BRL, Font(bold=True)
ws_pr[f"G{LIQ}"] = "Líquido por imóvel"
for j, col in enumerate(PCOLS):
    c = ws_pr[f"{col}{LIQ}"]
    c.value = f"=Painel!I{R1 + j}"
    c.number_format, c.font = BRL, LINK_FONT
box(ws_pr, f"A5:F{4 + NP}")
ws_pr.column_dimensions["G"].width = 18
for col in PCOLS:
    ws_pr.column_dimensions[col].width = 13

# ---------------------------------------------------------------- Repasses
title(ws_rp, "Repasses feitos", "Lance cada transferência feita ao proprietário.")
header(ws_rp, 4, ["Data", "Proprietário", "Valor (R$)", "Referência"], [12, 14, 13, 24])
for r, (d, nome, v, ref) in enumerate(repasses, 5):
    for col, val in enumerate([d, nome, v, ref], 1):
        ws_rp.cell(r, col, val).font = INPUT_FONT
    ws_rp.cell(r, 1).number_format = DATA
    ws_rp.cell(r, 3).number_format = BRL
box(ws_rp, f"A5:D{4 + len(repasses)}")

# ---------------------------------------------------------------- Extrato
title(ws_e, "Extrato do imóvel", "Escolha o imóvel na célula amarela. Pronto para imprimir em uma página A4.")
ws_e["A4"] = "Imóvel:"
ws_e["A4"].font = Font(bold=True)
ws_e["B4"] = "I05"
ws_e["B4"].fill, ws_e["B4"].font = YELLOW, Font(bold=True, color="0000FF")
dv = DataValidation(type="list", formula1=f"='Imóveis'!$A$5:$A${LI}", allow_blank=False)
ws_e.add_data_validation(dv)
dv.add("B4")
IMS = "'Imóveis'"
li = lambda col, sheet="Contratos", key="B": f"INDEX({sheet}!${col}$5:${col}${LI},MATCH($B$4,{sheet}!${key}$5:${key}${LI},0))"
info = [
    ("Endereço", "=" + li("B", IMS, "A") + '&" - "&' + li("C", IMS, "A") + '&", Curitiba - PR"', None),
    ("Contrato", f"={li('A')}", None),
    ("Inquilino", f"={li('C')}", None),
    ("Aluguel", f"={li('E')}", BRL),
    ("Garantia", f"={li('H')}", None),
    ("Próximo reajuste", f"={li('I')}", DATA),
    ("Índice", f"={li('G')}", None),
]
for i, (lab, f, fmt) in enumerate(info, 5):
    ws_e.cell(i, 1, lab)
    c = ws_e.cell(i, 2, f)
    c.font = LINK_FONT
    c.alignment = Alignment(horizontal="left")
    if fmt:
        c.number_format = fmt
header(ws_e, 13, ["Competência", "Vencimento", "Aluguel (R$)", "Multa (R$)", "Juros (R$)", "Recebido (R$)",
                  "Saldo (R$)", "Situação", "Faixa de atraso"], [36, 14, 14, 12, 11, 14, 14, 17, 14])
first = f"MATCH($B$4,{CR('C')},0)"
for k in range(6):
    r = 14 + k
    for col, src, fmt in [(1, "E", MES), (2, "F", DATA), (3, "G", BRL), (4, "K", BRL), (5, "L", BRL),
                          (6, "H", BRL), (7, "N", BRL), (8, "O", None), (9, "R", None)]:
        c = ws_e.cell(r, col, f"=INDEX({CR(src)},{first}+{k})")
        c.font = LINK_FONT
        if fmt:
            c.number_format = fmt
    ws_e.cell(r, 2).alignment = Alignment(horizontal="center")
ws_e["A20"] = "Total"
ws_e["A20"].font = Font(bold=True)
for col in "CDEFG":
    c = ws_e[f"{col}20"]
    c.value = f"=SUM({col}14:{col}19)"
    c.number_format, c.font = BRL, Font(bold=True)
box(ws_e, "A14:I20")
cores_situacao(ws_e, "H14:H19", "$H14")
summary = [
    ("Recebido (aluguel, multa e juros)", "=F20"),
    ("(-) Taxa de administração", f"=-ROUND(F20*{ADMC},2)"),
    ("(-) Despesas do proprietário", f"=-SUMIFS(Despesas!$D$5:$D$200,Despesas!$B$5:$B$200,$B$4,Despesas!$A$5:$A$200,\"<=\"&{REFC})"),
    ("(-) Custos da cobrança judicial", f"=-SUMIFS({JUD('E')},{JUD('B')},$B$4,{JUD('F')},\"Proprietário\",{JUD('A')},\"<=\"&{REFC})"),
    ("Líquido do proprietário", "=SUM(B22:B25)"),
]
for i, (lab, f) in enumerate(summary, 22):
    bold = i == 26
    ws_e.cell(i, 1, lab).font = Font(bold=bold)
    c = ws_e.cell(i, 2, f)
    c.number_format, c.font = BRL, Font(bold=bold)
    c.alignment = Alignment(horizontal="left")
ws_e["A28"] = "Divisão entre proprietários"
ws_e["A28"].font = Font(bold=True)
for k in range(NP):
    r = 29 + k
    ws_e.cell(r, 1, f"='Proprietários'!A{5 + k}").font = LINK_FONT
    c = ws_e.cell(r, 2, f"=INDEX('Proprietários'!${PC1}${5 + k}:${PCN}${5 + k},MATCH($B$4,'Proprietários'!${PC1}$4:${PCN}$4,0))")
    c.number_format, c.alignment = "0%", Alignment(horizontal="left")
    c = ws_e.cell(r, 3, f"=ROUND(B{r}*$B$26,2)")
    c.number_format = BRL
ws_e.page_setup.orientation = "landscape"
ws_e.page_setup.fitToWidth = 1
ws_e.page_setup.fitToHeight = 1
ws_e.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- Guia
title(ws_guia, "Controle de Aluguéis: guia rápido", "Exemplo com dados fictícios. Imóveis em Curitiba, contratos e valores inventados.")
guia = [
    ("Cores", "Texto azul = digitado. Texto verde = puxado de outra aba. Preto = fórmula. Fundo amarelo = parâmetro ou campo de escolha."),
    ("Rotina do dia a dia", "Ao receber um aluguel, lance uma linha em Recebimentos com a data, o código da cobrança (aba Cobranças) e o valor. O resto atualiza sozinho."),
    ("Situação da cobrança", "Calculada automaticamente: Pago, Pago com atraso, Parcial, Em atraso ou A vencer. Quando há negociação ou processo, escolha na coluna Tratativa (lista) \"Em acordo\" ou \"Cobrança judicial\"; enquanto houver saldo, a situação passa a mostrar a tratativa. Quando o saldo é quitado, volta a mostrar Pago com atraso."),
    ("Faixa de atraso", "Cada cobrança com saldo é classificada em 1 a 30, 31 a 60, 61 a 90 ou mais de 90 dias. O Painel mostra a quantidade e o saldo de cada faixa."),
    ("Cobrança judicial", "Lance custas, honorários e diligências na aba Cobrança judicial, indicando quem pagou e se o valor pode ser pedido ao inquilino na ação. Os custos pagos pelo proprietário são descontados no repasse. Confirme os valores reembolsáveis com o advogado responsável."),
    ("Multa e juros", "Multa de 10% sobre o aluguel e juros de mora de 1% ao mês pro rata die (0,0333% por dia), contados do vencimento até o pagamento ou até a data de referência. Os percentuais e a tolerância ficam na aba Parâmetros e devem seguir o contrato."),
    ("Correção monetária", "Não incluída neste exemplo. Pode ser acrescentada pelo índice do contrato (IGP-M ou IPCA)."),
    ("Taxa de administração", "10% sobre o valor recebido, descontada no repasse junto com as despesas e os custos judiciais pagos em nome do proprietário."),
    ("Repasses", "A aba Proprietários calcula o líquido de cada proprietário conforme a participação em cada imóvel. Lance cada transferência em Repasses e o saldo a repassar se atualiza. Saldo negativo indica que os custos superaram o recebido no período."),
    ("Reajuste anual", "A aba Contratos calcula a data do próximo reajuste pelo aniversário do contrato e destaca os que vencem em até 60 dias."),
    ("Extrato do imóvel", "Escolha o imóvel na célula amarela da aba Extrato: aparecem os meses, a faixa de atraso, o líquido e a divisão entre proprietários, prontos para imprimir ou enviar em PDF."),
    ("Data de referência", "Fixada em 28/10/2026 para o exemplo. No uso real, troque por =HOJE()."),
]
ws_guia.column_dimensions["A"].width = 24
ws_guia.column_dimensions["B"].width = 110
for i, (a, b) in enumerate(guia, 4):
    ws_guia.cell(i, 1, a).font = Font(bold=True)
    c = ws_guia.cell(i, 2, b)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws_guia.cell(i, 1).alignment = Alignment(vertical="top")

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.active = 1
wb.save(OUT)
print(f"{OUT.name}: {len(cobrancas)} cobranças, {len(recebimentos)} recebimentos, {len(repasses)} repasses")
